"""Build the Anatomn hackathon deck as a 16:9 PowerPoint.

Figma Slides imports this file as editable slides. The same deck is also
created natively by pitch/figma-plugin.
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

W = 1920
H = 1080
EMU = 13.333333 / 1920 * 914400

INK = "141816"
BONE = "F3ECDF"
BONE2 = "E7DFD0"
BLOOD = "9D2C2C"
MUTE = "8A8175"
MUTE2 = "A3988A"
LINE = "D9D0C3"
SRC_INK = "6F675E"
META_INK = "B7AD9F"
META_BLOOD = "F3D0CC"
CARD_LINE = "CFC6B8"

NOTES = [
    "Anatomn. АШУҮИС-ийн оюутнуудад зориулсан нэг вэб. Зузаан ном үүрэхгүй.",
    "1,000 оюутны асуудал нэг. Багш нар хүнд, зузаан номыг үнэтэй авдаг. Бид түүнийг нэг вэб апп дотор шийднэ. 1,000 нь эхний пилот, судалгааны хувь биш.",
    "Бодит үнэ. Gray's Anatomy for Students 5e, Elsevier, $94.99, 1,168 нүүр. Netter Atlas 9e, Elsevier original $84.99, 736 нүүр. Gray's Atlas, Matthews RUSH bookstore, $89.99. АНУ-ын нээлттэй үнэ. Монголд хүргэлт нэмэгдэнэ.",
    "$269.97 × Монголбанк 2026.10.28-ны ₮3,598.76 = 971,557₮, слайд дээр 972,000₮. Anatomn-ийн жилийн багц 49,000₮. 19.8 дахин, слайд дээр 20×.",
    "Нэг холбоос: 3D атлас, AI багш, шалгалт, хугарлын лаб, судсны лаб. AI багш Gray's-ийн 8 бүлгийн бүтцэд зангуудсан. Номын текстийг хуулдаггүй.",
    "Одоо нээлттэй. 2,763 сонгож болох бүтэц, 16 систем, насанд хүрэгчийн 205/206 яс, 6 шалгалт × 25 асуулт.",
    "Яагаад Anatomn. Монгол + латин. AI багш Gray's-ийн бүлэгт зангуудсан. Хугарлын лаб апп дотор. Ном дээр монгол хязгаартай, AI ба хугарал байхгүй. Гадаад 3D ихэвчлэн зөвхөн англи, хугарлын лаб өндөр үнэтэй, жил нь $80–120. Манай үнэ 49,000₮/жил.",
    "Сургууль авна. Нэг суудал жилд 49,000₮. 1,000 суудал = 49 сая₮. Оюутан 0₮. Сургуулиас гадуур сарын багц 9,900₮. Хоёр үнэ апп дээр байна. Төлбөр QPay, SocialPay, банкаар.",
    "1,000 оюутан × 972,000₮ ≈ 972 сая ном. Лиценз 49 сая. Зөрүү 923 сая. АШУҮИС 14,200 суралцагч, mnums.edu.mn, 2026.10.28. Их эмчийн кредит 170,000₮, Eguur 2025.08.20. ~30 кредит × 170,000 = 5.1 сая. 3 ном ≈ жилийн төлбөрийн 19%.",
    "Хүсэлт: АШУҮИС дээр 1,000 суудлын пилот. Холбоос нээлттэй.",
]


def emu(px):
    return Emu(int(round(px * EMU)))


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def blank(prs, bg):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(bg)
    return slide


def box(slide, x, y, w, h, color, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(x), emu(y), emu(w), emu(h))
    shape.shadow.inherit = False
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = emu(1)
    else:
        shape.line.fill.background()
    # Kill the default shape style line that python-pptx sometimes leaves.
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)
    return shape


def tracking(run, em):
    if run.font.size is None:
        return
    run._r.get_or_add_rPr().set("spc", str(int(round(em * run.font.size.pt * 100))))


def write(slide, value, x, y, w, h, size, color, weight=600, align="left", em=0, anchor="top"):
    shape = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE}[anchor]
    lines = value.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT}[align]
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        run.font.name = "Inter"
        run.font.size = Pt(size * 0.5)
        run.font.bold = weight >= 600
        run.font.color.rgb = rgb(color)
        tracking(run, em)
        # Medium weight is not a PPT boolean. Semi-bold rides on bold for 600+.
    return shape


def kicker(slide, label, dark):
    write(slide, label.upper(), 96, 64, 700, 36, 18, MUTE2 if dark else MUTE, 600, em=0.28)


def index(slide, n, dark):
    write(slide, n, 1600, 64, 224, 36, 18, MUTE2 if dark else MUTE, 500, align="right", em=0.18)


def source(slide, value, dark):
    write(slide, value, 96, 1000, 1728, 56, 16, SRC_INK if dark else MUTE, 400)


def deck():
    prs = Presentation()
    prs.slide_width = emu(W)
    prs.slide_height = emu(H)
    prs.core_properties.title = "Anatomn — Hackathon"
    prs.core_properties.subject = "АШУҮИС · 1,000 оюутан"

    s = blank(prs, INK)
    box(s, 0, 0, 14, H, BLOOD)
    write(s, "ANATOMN", 96, 64, 600, 40, 20, MUTE2, 600, em=0.42)
    index(s, "01", True)
    write(s, "Нэг вэб.\nЗузаан номгүй.", 96, 300, 1600, 320, 128, BONE, 600, em=-0.045)
    write(s, "АШУҮИС", 96, 960, 500, 48, 26, MUTE2, 500)
    write(s, "1,000 оюутан", 1200, 960, 624, 48, 26, MUTE2, 500, align="right")

    s = blank(prs, BONE)
    kicker(s, "Асуудал", False)
    index(s, "02", False)
    write(s, "1,000", 96, 300, 860, 240, 200, INK, 600, em=-0.06)
    write(s, "оюутан · нэг асуудал", 104, 560, 800, 60, 42, INK, 500)
    box(s, 980, 300, 1, 360, LINE)
    write(s, "Хүнд.\nЗузаан.\nҮнэтэй ном.", 1040, 320, 780, 320, 84, INK, 600, em=-0.045)
    source(s, "Багш, оюутан энэ номыг авдаг. Бүгдийг нэг вэб апп дотор.", False)

    s = blank(prs, BONE)
    kicker(s, "Бодит үнэ", False)
    index(s, "03", False)
    books = [
        (96, INK, BONE, META_INK, "$94.99", "Gray's Anatomy\nfor Students", "5-р хэвлэл · 1,168 нүүр", None),
        (681, BLOOD, BONE, META_BLOOD, "$84.99", "Netter Atlas\nof Human Anatomy", "9-р хэвлэл · 736 нүүр", None),
        (1266, BONE2, INK, MUTE, "$89.99", "Gray's Atlas\nof Anatomy", "Номын дэлгүүрийн үнэ", CARD_LINE),
    ]
    for x, bg, fg, meta, price, name, pages, line in books:
        box(s, x, 200, 557, 620, bg, line)
        write(s, price, x + 40, 236, 480, 90, 64, fg, 600, em=-0.04)
        write(s, name, x + 40, 620, 480, 100, 30, fg, 600)
        write(s, pages, x + 40, 740, 480, 40, 20, meta, 500)
    source(s, "Elsevier list · Matthews RUSH bookstore. АНУ-ын нээлттэй үнэ. Монголд хүргэлт нэмэгдэнэ.", False)

    s = blank(prs, INK)
    kicker(s, "Харьцуулалт", True)
    index(s, "04", True)
    write(s, "972,000₮", 96, 400, 720, 110, 80, BONE, 600, em=-0.05)
    write(s, "3 ном · нэг оюутан", 96, 520, 720, 44, 28, MUTE2, 500)
    write(s, "20×", 820, 384, 320, 130, 100, BLOOD, 600, em=-0.06)
    write(s, "49,000₮", 1180, 400, 640, 110, 80, BONE, 600, em=-0.05)
    write(s, "Anatomn · 1 жил", 1180, 520, 640, 44, 28, MUTE2, 500)
    source(s, "$269.97 × ₮3,598.76 = 971,557₮. Монголбанк, 2026.10.28. Жилийн үнэ бүтээгдэхүүн дээрх багц.", True)

    s = blank(prs, BONE)
    kicker(s, "Шийдэл", False)
    index(s, "05", False)
    write(s, "Нэг вэб апп.", 96, 150, 1400, 120, 84, INK, 600, em=-0.045)
    rows = ["3D атлас", "AI багш", "Шалгалт", "Хугарал", "Судас"]
    for i, name in enumerate(rows):
        y = 330 + i * 112
        box(s, 96, y, 1728, 1, LINE)
        write(s, f"0{i + 1}", 96, y + 32, 100, 48, 22, MUTE, 600, em=0.14)
        write(s, name, 216, y + 22, 1400, 64, 44, INK, 600, em=-0.03)
    box(s, 96, 330 + 5 * 112, 1728, 1, LINE)

    s = blank(prs, INK)
    kicker(s, "Одоо ажиллаж байна", True)
    index(s, "06", True)
    write(s, "2,763", 96, 200, 1400, 220, 180, BONE, 600, em=-0.06)
    write(s, "сонгож болох бүтэц", 96, 420, 800, 48, 32, MUTE, 500)
    for i, (num, label) in enumerate([("16", "биеийн систем"), ("205/206", "насанд хүрэгчийн яс"), ("6×25", "хугаацаатай шалгалт")]):
        x = 96 + i * 586
        write(s, num, x, 840, 520, 80, 60, BONE, 600, em=-0.04)
        write(s, label, x, 930, 520, 40, 22, MUTE, 500)

    s = blank(prs, INK)
    kicker(s, "Давуу тал", True)
    index(s, "07", True)
    write(s, "Яагаад Anatomn.", 96, 140, 1400, 100, 64, BONE, 600, em=-0.045)
    xs = [96, 620, 1080, 1500]
    for i, label in enumerate(["Anatomn", "Ном", "Гадаад 3D"]):
        write(s, label, xs[i + 1], 270, 380, 40, 22, "E7C3C0" if i == 0 else MUTE, 600)
    box(s, 96, 330, 1728, 1, "3A423D")
    rows = [
        ["Монгол хэл", "Монгол + Латин", "Хязгаартай", "Зөвхөн English"],
        ["AI багш", "Gray's дээр", "Байхгүй", "Байхгүй"],
        ["Хугарал", "Лаб дотор", "Байхгүй", "Өндөр үнэ"],
        ["Үнэ", "49,000₮ / жил", "Ном авна", "$80–120 / жил"],
    ]
    for r, row in enumerate(rows):
        y = 360 + r * 140
        for i, cell in enumerate(row):
            write(s, cell, xs[i], y + 28, 400, 56, 28, BONE if i < 2 else MUTE2, 600 if i < 2 else 500)
        box(s, 96, y + 120, 1728, 1, "2A302C")

    s = blank(prs, BONE)
    kicker(s, "Загвар", False)
    index(s, "08", False)
    write(s, "49,000₮", 96, 260, 1400, 170, 120, INK, 600, em=-0.055)
    write(s, "нэг суудал · нэг жил", 104, 450, 900, 56, 36, MUTE, 500)
    box(s, 0, 800, W, 280, INK)
    write(s, "49 сая₮", 96, 860, 900, 90, 64, BONE, 600, em=-0.04)
    write(s, "1,000 суудал · сургууль төлдөг", 96, 960, 900, 40, 22, MUTE2, 500)
    write(s, "Оюутан 0₮", 1100, 870, 724, 44, 28, BONE, 500, align="right")
    write(s, "Гадуурх багц 9,900₮ / сар", 900, 920, 924, 44, 24, MUTE2, 500, align="right")

    s = blank(prs, INK)
    kicker(s, "1,000 оюутан", True)
    index(s, "09", True)
    cols = [("972 сая", "номын төсөв", BONE), ("49 сая", "лиценз", BONE), ("923 сая", "зөрүү", BLOOD)]
    for i, (num, label, color) in enumerate(cols):
        x = 96 + i * 592
        write(s, num, x, 400, 540, 110, 76, color, 600, em=-0.05)
        write(s, label, x, 520, 540, 44, 28, MUTE2, 500)
    source(s, "АШУҮИС 14,200 суралцагч, 2026.10.28. Кредит 170,000₮ × ~30 = 5.1 сая, Eguur 2025.08.20. 3 ном ≈ жилийн төлбөрийн 19%.", True)

    s = blank(prs, BONE)
    kicker(s, "Хүсэлт", False)
    index(s, "10", False)
    write(s, "1,000 суудал.", 96, 280, 1600, 160, 108, INK, 600, em=-0.05)
    write(s, "anatomn-test.vercel.app/atlas", 96, 560, 1600, 70, 40, BLOOD, 500)
    write(s, "АШУҮИС · нэг холбоос · зузаан номгүй", 96, 960, 1400, 48, 24, MUTE, 500)

    for slide, note in zip(prs.slides, NOTES):
        slide.notes_slide.notes_text_frame.text = note
    prs.save("pitch/anatomn-hackathon.pptx")


if __name__ == "__main__":
    usd = 94.99 + 84.99 + 89.99
    mnt = usd * 3598.76
    assert round(usd, 2) == 269.97
    assert abs(mnt - 971557.24) < 1
    deck()
